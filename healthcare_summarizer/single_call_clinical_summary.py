import os
from dotenv import load_dotenv
import docx
import PyPDF2
from openai import OpenAI
from phi_masker import PHIMasker
import md_to_docx
from groq import Groq
from google import genai
from google.genai import types

load_dotenv()


def extract_text_from_docx(file_path):
    """Extract text from a DOCX file."""
    doc = docx.Document(file_path)
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    return '\n'.join(full_text)


def extract_text_from_pdf(file_path):
    """Extract text from a PDF file."""
    text = []
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            text.append(page.extract_text())
    return '\n'.join(text)


def extract_transcript(file_path):
    """Extract transcript text based on file extension."""
    file_extension = os.path.splitext(file_path)[1].lower()
    
    if file_extension == '.docx':
        return extract_text_from_docx(file_path)
    elif file_extension == '.pdf':
        return extract_text_from_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_extension}. Please use .docx or .pdf files.")


def generate_clinical_summary(transcript_text, use_masked_data=False):
    """PASS 1: Generate initial clinical summary without strict word limits."""
    
    system_content = (
    "You are a medical documentation assistant specialized in creating structured clinical summaries from doctor-patient conversation transcripts.\n"
    "This transcript contains masked PHI (Protected Health Information) for privacy protection. Conversations may be multilingual.\n\n"

    "Your responsibility is to extract and organize ONLY what is explicitly stated in the transcript. Do not infer, assume, extrapolate, or add context that is not clearly documented.\n\n"

    "## CORE PRINCIPLES\n"
    "Extract only documented information; never infer or extrapolate\n"
    "If information is not present, write \"Not documented\"\n"
    "Preserve clinical negations exactly (e.g., \"no fever\" must not become \"fever\")\n"
    "Resolve conflicting statements using the most recent clinician assessment\n"
    "Ensure that the document is complete and all the headings are present, even if some sections are \"Not documented\"\n"
    "Use standard medical terminology and abbreviations appropriately\n"
    "Exclude herbal teas and general supplements from Current Medications\n"
    "Current Medications must include ONLY prescription medications or explicitly named over-the-counter drugs with a defined dose or frequency."
    "Do NOT include herbal teas, botanical preparations, traditional remedies, supplements, nutraceuticals, or food-based products."
    "Examples that must NOT appear in Current Medications include (but are not limited to): herbal teas, calm tea, shatavari, ashwagandha, turmeric, supplements, powders, home remedies, or dietary items.\n"
    "Omit all patient and clinician identifying information\n"
    "Output clean Markdown only — no code fences, no explanations, no preamble\n\n"

    "Begin the response exactly with:\n"
    "# Clinical Summary\n\n"

    "## OUTPUT STRUCTURE\n\n"

    "# Clinical Summary\n\n"

    "## Quick Reference\n\n"

    "### Chief Complaint\n"
    "[Concise primary presenting concern]\n\n"

    "### Key Problems/Diagnoses\n"
    "[Bulleted list of active clinical issues]\n\n"

    "### Current Medications\n"
    "[Markdown table for multiple medications with dosing details; bulleted list for simple entries]\n\n"

    "### Follow-Up Actions\n"
    "[Numbered list for sequential tasks; bullets for non-sequential items]\n\n"

    "### Monitoring Needs\n"
    "[Parameters requiring ongoing surveillance]\n\n"

    "---\n\n"

    "## Brief Summary\n"
    "[Condensed overview using standard abbreviations - ADIME format with PES statement]\n\n"

    "**Assessment:**\n"
    "Essential patient background\n"
    "Key anthropometric data\n"
    "Critical dietary history\n"
    "Physical activity patterns\n"
    "Nutrition-focused physical findings\n"
    "Biochemical data (if documented)\n\n"

    "**Diagnosis:**\n"
    "[Primary nutrition diagnoses]\n"
    "**ICD-10 Code:** [Placeholder]\n"
    "**PES Statement:** [Problem] related to [Etiology] as evidenced by [Signs/Symptoms]\n\n"

    "**Intervention:**\n"
    "Nutrition prescription\n"
    "Food and nutrient delivery\n"
    "Patient education topics\n"
    "Counseling strategies\n"
    "Nutrition care coordination\n"
    "**SMART Goals:**\n"
    "- **S**pecific: [What exactly will be accomplished]\n"
    "- **M**easurable: [How progress will be quantified]\n"
    "- **A**chievable: [Realistic given patient's context]\n"
    "- **R**elevant: [Aligned with patient's needs]\n"
    "- **T**ime-bound: [Deadline or timeframe]\n\n"

    "**Monitoring and Evaluation:**\n"
    "Progress tracking\n"
    "Reviewing goals and outcomes\n"
    "Follow-up care plan\n\n"

    "---\n\n"

    "## Goldilocks Summary\n"
    "[Balanced clinical detail for standard documentation - ADIME format with PES statement]\n\n"

    "**Assessment:**\n"
    "Patient's personal history and reason for visit\n"
    "Anthropometric data and weight history\n"
    "Dietary history and patterns\n"
    "Physical activity patterns and limitations\n"
    "Nutrition-focused physical findings\n"
    "Biochemical data and lab results (if documented)\n"
    "Medications and supplements\n"
    "Lifestyle and social factors (occupation, sleep, stress)\n\n"

    "**Diagnosis:**\n"
    "[Working nutrition diagnoses with clinical reasoning]\n"
    "**ICD-10 Code:** [Placeholder]\n"
    "**PES Statement:** [Problem] related to [Etiology] as evidenced by [Signs/Symptoms]\n\n"

    "**Intervention:**\n"
    "Nutrition prescription with rationale\n"
    "Food and nutrient delivery strategies\n"
    "Education on specific nutrition guidelines\n"
    "Counseling strategies and motivational approaches\n"
    "Coordination of nutrition care\n"
    "**SMART Goals:**\n"
    "- **S**pecific: [What exactly will be accomplished]\n"
    "- **M**easurable: [How progress will be quantified]\n"
    "- **A**chievable: [Realistic given patient's context]\n"
    "- **R**elevant: [Aligned with patient's needs]\n"
    "- **T**ime-bound: [Deadline or timeframe]\n\n"

    "**Monitoring and Evaluation:**\n"
    "Progress evaluation metrics\n"
    "Reviewing goals and outcomes\n"
    "Follow-up care plan and timeline\n\n"

    "---\n\n"

    "## Detailed Summary\n"
    "[Comprehensive clinical narrative with context - ADIME format with PES statement]\n\n"

    "**Assessment:**\n"
    "Thorough patient background and reason for visit\n"
    "Anthropometric data with interpretation and weight history\n"
    "Gastrointestinal history (if documented)\n"
    "Detailed dietary patterns with examples and diet recall\n"
    "Diet history and previous diet attempts\n"
    "Any disordered eating thoughts or behaviors (if documented)\n"
    "Physical activity profile, history, limitations, and past injuries\n"
    "Exercise interests\n"
    "Nutrition-focused physical findings\n"
    "Biochemical data and medical test outcomes (if documented)\n"
    "Medications and supplements with dosing\n"
    "Lifestyle and social factors (occupation, sleep quality, stress levels)\n\n"

    "**Diagnosis:**\n"
    "[Comprehensive nutrition diagnoses with supporting evidence]\n"
    "**ICD-10 Code:** [Placeholder]\n"
    "**PES Statement:** [Problem] related to [Etiology] as evidenced by [Signs/Symptoms]\n\n"

    "**Intervention:**\n"
    "Detailed nutrition prescription with clinical rationale\n"
    "Food and nutrient delivery recommendations\n"
    "Education and counseling strategies with specific topics\n"
    "Coordination of nutrition care with other providers\n"
    "**SMART Goals:**\n"
    "- **S**pecific: [What exactly will be accomplished]\n"
    "- **M**easurable: [How progress will be quantified]\n"
    "- **A**chievable: [Realistic given patient's context]\n"
    "- **R**elevant: [Aligned with patient's needs]\n"
    "- **T**ime-bound: [Deadline or timeframe]\n\n"

    "**Monitoring and Evaluation:**\n"
    "Detailed progress tracking parameters\n"
    "Reviewing goals and outcomes\n"
    "Follow-up timeline and metrics\n\n"

    "---\n\n"

    "## Super Detailed Summary\n"
    "[Maximum clinical detail including patient quotes when documented - ADIME format with PES statement]\n\n"

    "**Assessment:**\n"
    "Exhaustive patient history and detailed reason for visit\n"
    "Complete anthropometric context with interpretation and comprehensive weight history\n"
    "Gastrointestinal history with symptoms (if documented)\n"
    "Extensive dietary history with symptoms, emotional responses, and specific examples\n"
    "Diet patterns, diet recall, and previous diet attempts\n"
    "Any disordered eating thoughts or behaviors (if documented)\n"
    "Physical activity history with preferences, limitations, past/existing injuries\n"
    "Exercise interests and barriers\n"
    "Nutrition-focused physical findings with clinical significance\n"
    "Comprehensive biochemical data and medical test outcomes (if documented)\n"
    "Complete medications and supplements list with dosing and frequency\n"
    "Lifestyle and social factors (occupation, sleep quality with patterns, stress levels)\n\n"

    "**Diagnosis:**\n"
    "[Complete diagnostic formulation with detailed clinical reasoning and evidence]\n"
    "**ICD-10 Code:** [Placeholder]\n"
    "**PES Statement:** [Problem] related to [Etiology] as evidenced by [Signs/Symptoms with specific data]\n\n"

    "**Intervention:**\n"
    "Comprehensive nutrition prescription with detailed rationale\n"
    "Food and nutrient delivery with specific strategies\n"
    "Education topics with key teaching points\n"
    "Counseling strategies with motivational interviewing techniques\n"
    "Nutrition care coordination with interdisciplinary team\n"
    "**SMART Goals:**\n"
    "- **S**pecific: [What exactly will be accomplished with details]\n"
    "- **M**easurable: [How progress will be quantified with specific metrics]\n"
    "- **A**chievable: [Realistic given patient's context with supporting rationale]\n"
    "- **R**elevant: [Aligned with patient's needs and priorities]\n"
    "- **T**ime-bound: [Deadline or timeframe with follow-up schedule]\n\n"

    "**Monitoring and Evaluation:**\n"
    "Detailed progress evaluation with specific parameters\n"
    "Reviewing goals and outcomes with benchmarks\n"
    "Follow-up metrics, timeline, and contingency plans\n\n"

    "---\n\n"

    "## SOAP Note\n"
    "[Standard SOAP format]\n\n"

    "**Client/Patient Name:** <confidential>\n"
    "**Assessment Date:** <confidential>\n"
    "**Submitted By:** <confidential>\n\n"

    "### Subjective\n"
    "[Patient-reported symptoms, concerns, and relevant history]\n\n"

    "### Objective\n"
    "[Observable findings, measurements, and vital signs when documented]\n\n"

    "### Assessment\n"
    "[Clinical impressions and diagnostic reasoning]\n\n"

    "### Plan\n"
    "[Bulleted action items with interventions and follow-up]\n\n"

    "---\n\n"

    "## MARKDOWN FORMATTING STANDARDS\n"
    "Headers: # (title), ## (main sections), ### (subsections)\n"
    "Tables: Use Markdown tables for structured clinical data (medications, vitals, labs)\n"
    "Lists: Use - for bullets and 1., 2., 3. for numbered lists (restart numbering per section)\n"
    "Emphasis: Use **bold** for important findings and italic for patient quotes\n"
    "Spacing: Leave blank lines between sections for readability\n"
    "PES Statement Format: [Problem] related to [Etiology] as evidenced by [Signs/Symptoms]\n"
    "SMART Goals Format: Explicitly label each component (Specific, Measurable, Achievable, Relevant, Time-bound)\n\n"

    "Return only the final Markdown output with complete information extracted from the transcript."
)
    
    client = genai.Client()

    # Create user message
    user_message = types.Content(
        role="user",
        parts=[types.Part.from_text(text=f"Please create a structured clinical summary from the following doctor-patient transcript:\n\n{transcript_text}")]
    )

    contents = [user_message]

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="text/plain",
            system_instruction=system_content,
            temperature=0.1,
        )
    )

    return response.text if response else ""


def condense_clinical_summary(initial_summary):
    """PASS 2: Condense the summary to meet strict word count requirements."""
    
    condensing_prompt = (
    "You are a clinical documentation editor.\n\n"
    
    "You will be given a clinical summary that already follows the required template with ADIME format and PES statements.\n\n"
    
    "Your task is NOT to add new information.\n"
    "Your task is to EDIT and CONDENSE the following sections so they strictly comply with word limits:\n\n"
    
    "- **Brief Summary**: 150–300 words\n"
    "- **Goldilocks Summary**: 200–350 words\n"
    "- **Detailed Summary**: 300–400 words\n"
    "- **Super Detailed Summary**: 400–500 words\n"
    "- **SOAP Note** (entire SOAP section combined): 200–300 words\n\n"
    
    "## STRICT RULES:\n"
    "- Do NOT remove clinically important facts\n"
    "- Do NOT add new content or infer information\n"
    "- Remove redundancy and compress phrasing where possible\n"
    "- Preserve headings and structure EXACTLY as provided\n"
    "- Maintain ADIME format (Assessment, Diagnosis, Intervention, Monitoring/Evaluation)\n"
    "- Keep PES statements in standard format: [Problem] related to [Etiology] as evidenced by [Signs/Symptoms]\n"
    "- Preserve SMART Goals format with explicit labels (Specific, Measurable, Achievable, Relevant, Time-bound)\n"
    "- Keep medical terminology intact\n"
    "- Preserve clinical negations (e.g., 'no fever' must remain 'no fever')\n"
    "- Keep ICD-10 Code placeholders\n"
    "- Ensure final word counts fall strictly within the specified ranges\n"
    "- Keep the Quick Reference section unchanged\n\n"
    
    "Return ONLY the revised Markdown with no code fences, no preamble, no explanations."
)
    
    client = genai.Client()

    # Create user message with the initial summary
    user_message = types.Content(
        role="user",
        parts=[types.Part.from_text(text=f"Please condense the following clinical summary to meet the word count requirements:\n\n{initial_summary}")]
    )

    contents = [user_message]

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="text/plain",
            system_instruction=condensing_prompt,
            temperature=0.1,
        )
    )

    return response.text if response else ""


if __name__ == "__main__":
    # Configuration
    input_file = "transcript.pdf"
    output_md_file = "clinical_summary.md"
    output_docx_file = "clinical_summary.docx"
    masked_transcript_file = "masked_transcript.txt"
    
    try:
        # Extract and mask transcript
        print(f"Extracting transcript from {input_file}...")
        transcript = extract_transcript(input_file)

        print("Applying PHI masking...")
        masker = PHIMasker(use_consistent_hashing=True)
        masked_transcript = masker.apply_all_masking(transcript)
        
        # Save masked transcript
        with open(masked_transcript_file, 'w', encoding='utf-8') as f:
            f.write(masked_transcript)
        f.close()
        
        # PASS 1: Generate initial Markdown summary
        print("PASS 1: Generating initial clinical summary...")
        initial_summary = generate_clinical_summary(masked_transcript, use_masked_data=True)
        MAX_RETRIES = 5
        retries = 0
        while len(initial_summary) == 0 and retries < MAX_RETRIES:
            print("Received empty summary, retrying...")
            initial_summary = generate_clinical_summary(masked_transcript, use_masked_data=True)
            retries += 1

        if retries == MAX_RETRIES:
            print("Reached maximum retry limit for initial generation, please try again")
            exit(1)
        
        # Clean up response (remove code fences if LLM added them)
        initial_summary = initial_summary.strip()
        if initial_summary.startswith('```'):
            initial_summary = initial_summary.replace('```markdown', '').replace('```', '').strip()
        
        # PASS 2: Condense to meet word count requirements
        print("PASS 2: Condensing summary to meet word count requirements...")
        condensed_summary = condense_clinical_summary(initial_summary)
        retries = 0
        while len(condensed_summary) == 0 and retries < MAX_RETRIES:
            print("Received empty condensed summary, retrying...")
            condensed_summary = condense_clinical_summary(initial_summary)
            retries += 1

        if retries == MAX_RETRIES:
            print("Reached maximum retry limit for condensing, using initial summary")
            markdown_summary = initial_summary
        else:
            # Clean up condensed response
            condensed_summary = condensed_summary.strip()
            if condensed_summary.startswith('```'):
                condensed_summary = condensed_summary.replace('```markdown', '').replace('```', '').strip()
            markdown_summary = condensed_summary
        
        # Save Markdown
        with open(output_md_file, 'w', encoding='utf-8') as f:
            f.write(markdown_summary)
        print(f"✓ Markdown summary saved: {output_md_file}")
        
        # Print summary
        print("\n" + "="*80)
        print("CLINICAL SUMMARY (MARKDOWN)")
        print("="*80 + "\n")
        print(markdown_summary)
        print("\n" + "="*80 + "\n")
        
        # Convert Markdown to DOCX
        print("Converting Markdown to DOCX...")
        md_to_docx.markdown_to_docx(markdown_summary, output_docx_file)
        print(f"✓ DOCX file created successfully: {output_docx_file}")
        
    except Exception as e:
        print(f"An error occurred: {e}")
        import traceback
        traceback.print_exc()
